"""Marp Markdown sunumunu düzenlenebilir, küçük boyutlu bir PPTX'e çevirir (python-pptx).

Marp'ın PPTX çıktısı her slaytı resim olarak gömer (büyük ve düzenlenemez). Bu modül başlık, liste, tablo, kod,
alıntı ve konuşma notlarını PowerPoint'in kendi nesneleriyle yazar. Slayt sınıfları: baslik, bolum, yogun, sema.
"""
import math
import os
import pathlib
from PIL import Image
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

TEAL = RGBColor(0x00, 0x6D, 0x77)
KOYU = RGBColor(0x1F, 0x2D, 0x3A)
GRI = RGBColor(0x5B, 0x67, 0x70)
ACIK = RGBColor(0xEE, 0xF7, 0xF8)
KOD_ZEMIN = RGBColor(0xF3, 0xF4, 0xF6)
BEYAZ = RGBColor(0xFF, 0xFF, 0xFF)

G, Y = Inches(13.333), Inches(7.5)
SOL, SAG = Inches(0.7), Inches(0.7)
GOVDE_UST, GOVDE_ALT = Inches(1.45), Inches(6.85)
GENIS = G - SOL - SAG


# ------------------------------------------------------------------ ayrıştırma
def on_bilgi(md):
    m = re.match(r'\A---\n(.*?)\n---\n', md, flags=re.S)
    if not m:
        return {}, md
    bilgi = {}
    for satir in m.group(1).split('\n'):
        k = re.match(r'^(\w+):\s*"?(.*?)"?\s*$', satir)
        if k:
            bilgi[k.group(1)] = k.group(2)
    return bilgi, md[m.end():]


def slaytlar(md):
    bolumler, simdiki, kod = [], [], False
    for satir in md.split('\n'):
        if satir.startswith('```'):
            kod = not kod
        if satir.strip() == '---' and not kod:
            bolumler.append('\n'.join(simdiki))
            simdiki = []
        else:
            simdiki.append(satir)
    bolumler.append('\n'.join(simdiki))
    return [b for b in bolumler if b.strip()]


def cozumle(metin):
    sinif, notlar = '', []
    for yorum in re.findall(r'<!--(.*?)-->', metin, flags=re.S):
        y = yorum.strip()
        m = re.match(r'_class:\s*(\w+)', y)
        if m:
            sinif = m.group(1)
        elif not re.match(r'_?\w+:', y):
            notlar.append(re.sub(r'^Konuşma notu:\s*', '', y))
    metin = re.sub(r'<!--.*?-->', '', metin, flags=re.S)
    bloklar, satirlar, i = [], metin.split('\n'), 0
    while i < len(satirlar):
        s = satirlar[i]
        if not s.strip():
            i += 1
            continue
        if s.startswith('```'):
            dil, kod, i = s[3:].strip(), [], i + 1
            while i < len(satirlar) and not satirlar[i].startswith('```'):
                kod.append(satirlar[i])
                i += 1
            bloklar.append(('kod', dil, kod))
            i += 1
            continue
        m = re.match(r'^(#{1,3})\s+(.*)', s)
        if m:
            bloklar.append((f'h{len(m.group(1))}', m.group(2)))
            i += 1
            continue
        if s.startswith('|'):
            satir = []
            while i < len(satirlar) and satirlar[i].startswith('|'):
                hucre = [h.strip() for h in satirlar[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', h) for h in hucre if h):
                    satir.append(hucre)
                i += 1
            bloklar.append(('tablo', satir))
            continue
        if s.lstrip().startswith('>'):
            alinti = []
            while i < len(satirlar) and satirlar[i].lstrip().startswith('>'):
                alinti.append(satirlar[i].lstrip()[1:].strip())
                i += 1
            bloklar.append(('alinti', ' '.join(a for a in alinti if a)))
            continue
        if re.match(r'^\s*([-*]|\d+\.)\s+', s):
            ogeler = []
            while i < len(satirlar) and (re.match(r'^\s*([-*]|\d+\.)\s+', satirlar[i]) or
                                         (satirlar[i].startswith('   ') and satirlar[i].strip() and ogeler)):
                m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)', satirlar[i])
                if m:
                    ogeler.append([len(m.group(1)) // 2, m.group(2)[0].isdigit(), m.group(3)])
                else:
                    ogeler[-1][2] += ' ' + satirlar[i].strip()
                i += 1
            bloklar.append(('liste', ogeler))
            continue
        mg = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', s.strip())
        if mg:
            bloklar.append(('gorsel', mg.group(1), mg.group(2)))
            i += 1
            continue
        paragraf = []
        while i < len(satirlar) and satirlar[i].strip() and not re.match(r'^(```|#|\||>|!\[|\s*([-*]|\d+\.)\s)', satirlar[i]):
            paragraf.append(satirlar[i].strip())
            i += 1
        bloklar.append(('p', ' '.join(paragraf)))
    return sinif, bloklar, '\n\n'.join(notlar)


# ------------------------------------------------------------------ metin yazma
SATIRICI = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))')


def temizle(metin):
    return metin.replace('\\+', '+').replace('\\', '')


def calistir_yaz(paragraf, metin, boyut, renk=KOYU, kalin=False):
    for parca in SATIRICI.split(temizle(metin)):
        if not parca:
            continue
        run = paragraf.add_run()
        run.font.size = Pt(boyut)
        run.font.color.rgb = renk
        run.font.bold = kalin
        if parca.startswith('**') and parca.endswith('**'):
            run.text, run.font.bold = parca[2:-2], True
        elif parca.startswith('`') and parca.endswith('`'):
            run.text = parca[1:-1]
            run.font.name = 'Consolas'
            run.font.color.rgb = RGBColor(0xA1, 0x12, 0x3F) if renk == KOYU else renk
        elif parca.startswith('*') and parca.endswith('*') and len(parca) > 2:
            run.text, run.font.italic = parca[1:-1], True
        elif parca.startswith('['):
            run.text = re.match(r'\[([^\]]+)\]', parca).group(1)
        else:
            run.text = parca


# ------------------------------------------------------------------ görseller
KAYNAK_KLASOR = None


def gorsel_yolu(yol):
    """Markdown'daki göreli yolu gerçek dosyaya çevirir. PPTX SVG alamadığı için
    PNG kardeşi tercih edilir. Slayt kaynağı slides/week-N altında, görseller ise
    docs/week-N/assets altında olduğu için oraya da bakılır."""
    if KAYNAK_KLASOR is None:
        return None
    kok = pathlib.Path(KAYNAK_KLASOR)
    adaylar = [kok / yol, kok.parent.parent / 'docs' / kok.name / yol]
    for a in adaylar:
        png = a.with_suffix('.png')
        if png.exists():
            return png
        if a.exists() and a.suffix.lower() != '.svg':
            return a
    return None


def gorsel_olcu(blok, genis_in):
    """(genişlik_in, yükseklik_in) — Marp 'w:NNN' ipucunu da dikkate alır."""
    _, alt, yol = blok
    p = gorsel_yolu(yol)
    if not p:
        return 0.0, 0.0
    with Image.open(p) as im:
        gw, gh = im.size
    oran = gh / float(gw)
    g = genis_in
    m = re.search(r'w:(\d+)', alt or '')
    if m:
        g = min(genis_in, genis_in * (int(m.group(1)) / 1280.0))
    return g, g * oran


def yaz_gorsel(slayt, blok, ust, boyut, kalan_in):
    g, y = gorsel_olcu(blok, GENIS / 914400)
    if g <= 0:
        return 0.0
    if kalan_in and y > kalan_in:
        oran = kalan_in / y
        g, y = g * oran, kalan_in
    p = gorsel_yolu(blok[2])
    sol = SOL + int((GENIS - Inches(g)) / 2)
    slayt.shapes.add_picture(str(p), sol, ust, width=Inches(g), height=Inches(y))
    return y


def metin_kutusu(slayt, sol, ust, gen, yuk):
    kutu = slayt.shapes.add_textbox(sol, ust, gen, yuk)
    tf = kutu.text_frame
    tf.word_wrap = True
    for kenar in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, kenar, Inches(0.05))
    return kutu, tf


# ------------------------------------------------------------------ yükseklik tahmini
def satir_sayisi(metin, boyut, genislik_in, cift=1.0):
    karakter = max(len(re.sub(r'[*`]', '', temizle(metin))), 1)
    karakter_gen = boyut * 0.52 / 72 * cift          # ortalama karakter genişliği (inç)
    return max(1, math.ceil(karakter * karakter_gen / genislik_in))


def blok_yuksekligi(blok, boyut, genislik_in):
    tur = blok[0]
    sat = boyut * 1.3 / 72
    if tur == 'p':
        return satir_sayisi(blok[1], boyut, genislik_in) * sat + 0.12
    if tur == 'alinti':
        return satir_sayisi(blok[1], boyut, genislik_in - 0.4) * sat + 0.3
    if tur == 'liste':
        return sum(satir_sayisi(o[2], boyut, genislik_in - 0.4 - 0.35 * o[0]) * sat + 0.05 for o in blok[1]) + 0.1
    if tur == 'kod':
        kod_boyut = max(boyut * 0.62, 10)
        return len(blok[2]) * kod_boyut * 1.2 / 72 + 0.3
    if tur == 'tablo':
        tb = max(boyut * 0.72, 10)
        sutun = max(len(r) for r in blok[1])
        toplam = 0
        for r in blok[1]:
            en_uzun = max(satir_sayisi(h, tb, genislik_in / sutun) for h in r) if r else 1
            toplam += en_uzun * tb * 1.3 / 72 + 0.12
        return toplam + 0.15
    if tur == 'gorsel':
        return gorsel_olcu(blok, genislik_in)[1]
    if tur in ('h2', 'h3'):
        return boyut * 1.5 / 72 + 0.1
    return 0.3


# ------------------------------------------------------------------ bloklar
def yaz_liste(slayt, blok, ust, boyut):
    yuk = blok_yuksekligi(blok, boyut, GENIS / 914400)
    _, tf = metin_kutusu(slayt, SOL, ust, GENIS, Inches(yuk))
    numara = {}
    for j, (duzey, sirali, metin) in enumerate(blok[1]):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.level = min(duzey, 4)
        p.space_after = Pt(3)
        if sirali:
            numara[duzey] = numara.get(duzey, 0) + 1
            isaret = f'{numara[duzey]}. '
        else:
            isaret = '• ' if duzey == 0 else '– '
        run = p.add_run()
        run.text = '    ' * duzey + isaret
        run.font.size, run.font.color.rgb, run.font.bold = Pt(boyut), TEAL, True
        calistir_yaz(p, metin, boyut)
    return yuk


def yaz_tablo(slayt, blok, ust, boyut):
    satirlar = blok[1]
    sutun = max(len(r) for r in satirlar)
    satirlar = [r + [''] * (sutun - len(r)) for r in satirlar]
    tb = max(boyut * 0.72, 10)
    yuk = blok_yuksekligi(blok, boyut, GENIS / 914400)
    uzunluk = [max(len(temizle(r[c])) for r in satirlar) + 4 for c in range(sutun)]
    toplam_uz = sum(uzunluk)
    gen_tablo = GENIS if toplam_uz > 60 else int(GENIS * max(0.55, toplam_uz / 60))
    sekil = slayt.shapes.add_table(len(satirlar), sutun, SOL, ust, gen_tablo, Inches(yuk))
    tablo = sekil.table
    for c in range(sutun):
        tablo.columns[c].width = int(gen_tablo * uzunluk[c] / toplam_uz)
    for r, satir in enumerate(satirlar):
        for c, hucre in enumerate(satir):
            h = tablo.cell(r, c)
            h.margin_left = h.margin_right = Inches(0.06)
            h.margin_top = h.margin_bottom = Inches(0.03)
            h.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = h.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if r == 0:
                h.fill.solid()
                h.fill.fore_color.rgb = TEAL
                calistir_yaz(p, hucre, tb, BEYAZ, kalin=True)
            else:
                h.fill.solid()
                h.fill.fore_color.rgb = ACIK if r % 2 == 0 else BEYAZ
                calistir_yaz(p, hucre, tb)
    return yuk


def yaz_kod(slayt, blok, ust, boyut):
    kod_boyut = max(boyut * 0.62, 10)
    yuk = blok_yuksekligi(blok, boyut, GENIS / 914400)
    kutu = slayt.shapes.add_shape(MSO_SHAPE.RECTANGLE, SOL, ust, GENIS, Inches(yuk))
    kutu.fill.solid()
    kutu.fill.fore_color.rgb = KOD_ZEMIN
    kutu.line.color.rgb = RGBColor(0xD9, 0xDE, 0xE3)
    tf = kutu.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = tf.margin_bottom = Inches(0.1)
    for j, satir in enumerate(blok[2] or ['']):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = satir if satir else ' '
        run.font.name = 'Consolas'
        run.font.size = Pt(kod_boyut)
        run.font.color.rgb = KOYU
    return yuk


def yaz_paragraf(slayt, metin, ust, boyut, renk=KOYU, hizala=PP_ALIGN.LEFT, alinti=False):
    yuk = blok_yuksekligi(('alinti' if alinti else 'p', metin), boyut, GENIS / 914400)
    if alinti:
        kutu = slayt.shapes.add_shape(MSO_SHAPE.RECTANGLE, SOL, ust, GENIS, Inches(yuk))
        kutu.fill.solid()
        kutu.fill.fore_color.rgb = ACIK
        kutu.line.color.rgb = TEAL
        tf = kutu.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
    else:
        _, tf = metin_kutusu(slayt, SOL, ust, GENIS, Inches(yuk))
    p = tf.paragraphs[0]
    p.alignment = hizala
    calistir_yaz(p, metin, boyut, renk)
    return yuk


# ------------------------------------------------------------------ slayt çerçevesi
def cerceve(slayt, bilgi, no, logo, koyu=False):
    renk = RGBColor(0xD8, 0xF0, 0xF2) if koyu else GRI
    if logo and not koyu:
        slayt.shapes.add_picture(logo, SOL, Inches(0.18), height=Inches(0.36))
    _, tf = metin_kutusu(slayt, SOL + Inches(0.45), Inches(0.2), Inches(8), Inches(0.35))
    calistir_yaz(tf.paragraphs[0], bilgi.get('header', ''), 11, renk)
    _, tf = metin_kutusu(slayt, SOL, Y - Inches(0.45), Inches(8), Inches(0.3))
    calistir_yaz(tf.paragraphs[0], bilgi.get('footer', ''), 10, renk)
    if no:
        _, tf = metin_kutusu(slayt, G - SAG - Inches(1), Y - Inches(0.45), Inches(1), Inches(0.3))
        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        calistir_yaz(tf.paragraphs[0], str(no), 10, renk)


def donustur(md_yolu, pptx_yolu, logo=None):
    global KAYNAK_KLASOR
    KAYNAK_KLASOR = os.path.dirname(os.path.abspath(md_yolu))
    md = open(md_yolu, encoding='utf-8').read()
    bilgi, govde = on_bilgi(md)
    sunum = Presentation()
    sunum.slide_width, sunum.slide_height = G, Y
    bos = sunum.slide_layouts[6]
    for no, ham in enumerate(slaytlar(govde), start=1):
        sinif, bloklar, notlar = cozumle(ham)
        slayt = sunum.slides.add_slide(bos)
        koyu = sinif == 'bolum'
        if koyu:
            arka = slayt.background.fill
            arka.solid()
            arka.fore_color.rgb = TEAL
        cerceve(slayt, bilgi, None if no == 1 else no, logo, koyu)

        if sinif in ('baslik', 'bolum'):
            baslik = next((b[1] for b in bloklar if b[0] == 'h1'), '')
            digerleri = [b for b in bloklar if b[0] != 'h1']
            ust = Inches(2.4) if digerleri else Inches(3.0)
            _, tf = metin_kutusu(slayt, SOL, ust, GENIS, Inches(1.4))
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER if sinif == 'baslik' else PP_ALIGN.LEFT
            calistir_yaz(tf.paragraphs[0], baslik, 40, BEYAZ if koyu else TEAL, kalin=True)
            ust += Inches(1.5)
            for b in digerleri:
                if b[0] == 'p':
                    ust += Inches(yaz_paragraf(slayt, b[1], ust, 20, BEYAZ if koyu else KOYU,
                                               PP_ALIGN.CENTER if sinif == 'baslik' else PP_ALIGN.LEFT) + 0.05)
        else:
            baslik = next((b[1] for b in bloklar if b[0] == 'h1'), '')
            if baslik:
                _, tf = metin_kutusu(slayt, SOL, Inches(0.6), GENIS, Inches(0.8))
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                calistir_yaz(tf.paragraphs[0], baslik, 30 if sinif != 'yogun' else 26, TEAL, kalin=True)
            govde_bloklar = [b for b in bloklar if b[0] != 'h1']
            boyut = 18 if sinif == 'yogun' else 20
            alan = (GOVDE_ALT - GOVDE_UST) / 914400
            while boyut > 11:
                toplam = sum(blok_yuksekligi(b, boyut, GENIS / 914400) + 0.12 for b in govde_bloklar)
                if toplam <= alan:
                    break
                boyut -= 1
            ust = GOVDE_UST
            for b in govde_bloklar:
                if b[0] == 'liste':
                    yuk = yaz_liste(slayt, b, ust, boyut)
                elif b[0] == 'tablo':
                    yuk = yaz_tablo(slayt, b, ust, boyut)
                elif b[0] == 'kod':
                    yuk = yaz_kod(slayt, b, ust, boyut)
                elif b[0] == 'alinti':
                    yuk = yaz_paragraf(slayt, b[1], ust, boyut, alinti=True)
                elif b[0] == 'gorsel':
                    yuk = yaz_gorsel(slayt, b, ust, boyut, (GOVDE_ALT - ust) / 914400)
                elif b[0] in ('h2', 'h3'):
                    yuk = yaz_paragraf(slayt, f'**{b[1]}**', ust, boyut + 2, TEAL)
                else:
                    yuk = yaz_paragraf(slayt, b[1], ust, boyut)
                ust += Inches(yuk + 0.12)
        if notlar:
            slayt.notes_slide.notes_text_frame.text = notlar
    sunum.core_properties.title = bilgi.get('title', '')
    sunum.core_properties.author = bilgi.get('author', '')
    sunum.save(pptx_yolu)
    return no
